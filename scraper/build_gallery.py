import json,os,hashlib,re,collections
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,PatternFill,Border,Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

items=json.load(open('items.json'))
MARKUP=2.10
def key(u): return hashlib.md5(u.encode()).hexdigest()
def short(t,n):
    t=re.sub(r'\s+',' ',t or '').strip()
    return t if len(t)<=n else t[:n-1].rsplit(' ',1)[0]+'…'

items.sort(key=lambda x:(x['source'],x['kind'],x['price']))
MAXP=max(len(i['gallery']) for i in items)

wb=Workbook(); ws=wb.active; ws.title='All Photos'
NAVY='FF1F3864'; ACCENT='FFD9E2F3'; ZEBRA='FFF7F9FC'; GREEN='FFE2EFDA'
thin=Side(style='thin',color='FFBFBFBF'); border=Border(left=thin,right=thin,top=thin,bottom=thin)

INFO=[('#',6),('Product Name',40),('Description',52),('Weight',18),
      ('Metal',19),('Price (USD)',13),('FINAL PRICE (+110%)',19),('Source',16),('Photos',9)]
NI=len(INFO)
TOTAL=NI+MAXP

ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=TOTAL)
t=ws.cell(1,1,'LAB DIAMOND JEWELRY — 18K GOLD & 925 SILVER — FULL PHOTO GALLERY')
t.font=Font(size=16,bold=True,color='FFFFFFFF'); t.fill=PatternFill('solid',fgColor=NAVY)
t.alignment=Alignment(horizontal='center',vertical='center'); ws.row_dimensions[1].height=30
ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=TOTAL)
s=ws.cell(2,1,'Every photo published for each piece · FINAL PRICE = Price × 2.10 (cost + 110%) · Prices in USD · No links in this file — see the main workbook')
s.font=Font(size=10,italic=True,color=NAVY); s.fill=PatternFill('solid',fgColor=ACCENT)
s.alignment=Alignment(horizontal='center',vertical='center'); ws.row_dimensions[2].height=18

HR=3
for c,(n,w) in enumerate(INFO,1):
    cell=ws.cell(HR,c,n); cell.font=Font(bold=True,color='FFFFFFFF',size=10)
    cell.fill=PatternFill('solid',fgColor=NAVY); cell.border=border
    cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.column_dimensions[get_column_letter(c)].width=w
for k in range(MAXP):
    c=NI+1+k
    cell=ws.cell(HR,c,f'Photo {k+1}')
    cell.font=Font(bold=True,color='FFFFFFFF',size=9)
    cell.fill=PatternFill('solid',fgColor='FF2E5C9A'); cell.border=border
    cell.alignment=Alignment(horizontal='center',vertical='center')
    ws.column_dimensions[get_column_letter(c)].width=11.6
ws.row_dimensions[HR].height=30

IMG=78
r=HR+1; missing=0; placed=0
for n,it in enumerate(items,1):
    ws.row_dimensions[r].height=62
    vals=[n,it['title'],short(it['desc'],500),it['weight_display'],
          it['metal_detail'] or it['metal'],round(it['price'],2),None,
          it['source'],len(it['gallery'])]
    for c,v in enumerate(vals,1):
        cell=ws.cell(r,c,v); cell.border=border
        cell.alignment=Alignment(vertical='center',wrap_text=(c in (2,3)),
            horizontal='center' if c in (1,4,5,8,9) else 'left')
        if r%2==0: cell.fill=PatternFill('solid',fgColor=ZEBRA)
    ws.cell(r,6).number_format='"$"#,##0.00'
    f=ws.cell(r,7,f'=F{r}*{MARKUP}'); f.number_format='"$"#,##0.00'
    f.font=Font(bold=True,color='FF006100'); f.fill=PatternFill('solid',fgColor=GREEN)
    f.border=border; f.alignment=Alignment(horizontal='center',vertical='center')
    for k,u in enumerate(it['gallery']):
        c=NI+1+k
        cell=ws.cell(r,c); cell.border=border
        if r%2==0: cell.fill=PatternFill('solid',fgColor=ZEBRA)
        p='gal2/'+key(u)+'.jpg'
        if os.path.exists(p):
            im=XLImage(p); im.width=im.height=IMG
            im.anchor=f'{get_column_letter(c)}{r}'; ws.add_image(im); placed+=1
        else: missing+=1
    r+=1

ws.freeze_panes=f'{get_column_letter(NI+1)}{HR+1}'
ws.auto_filter.ref=f'A{HR}:{get_column_letter(NI)}{r-1}'
ws.sheet_view.showGridLines=False

# ---- summary ----
s2=wb.create_sheet('Summary'); s2.sheet_view.showGridLines=False
s2.column_dimensions['A'].width=34; s2.column_dimensions['B'].width=76
def hdr(row,txt):
    s2.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2)
    c=s2.cell(row,1,txt); c.font=Font(bold=True,size=12,color='FFFFFFFF')
    c.fill=PatternFill('solid',fgColor=NAVY); s2.row_dimensions[row].height=22
def kv(row,k,v):
    a=s2.cell(row,1,k); a.font=Font(bold=True,size=10); a.alignment=Alignment(vertical='top')
    b=s2.cell(row,2,str(v)); b.alignment=Alignment(wrap_text=True,vertical='top')
row=1
hdr(row,'FULL PHOTO GALLERY — SUMMARY'); row+=2
kv(row,'Unique items',len(items)); row+=1
kv(row,'Total photos embedded',placed); row+=1
kv(row,'Photos per item','average %.1f · maximum %d'%(sum(len(i['gallery']) for i in items)/len(items),MAXP)); row+=1
kv(row,'Markup','FINAL PRICE = Price × 2.10  (original price + 110%)'); row+=1
kv(row,'Links','Deliberately omitted from this file. Product and image links are in the main workbook.'); row+=2
hdr(row,'SOURCES'); row+=1
for k,v in collections.Counter(i['source'] for i in items).most_common(): kv(row,k,f'{v} items'); row+=1
kv(row,'StarsGem (excluded)','Quote-only site with no published prices.')
s2.cell(row,2).font=Font(color='FFC00000',size=10); row+=2
hdr(row,'BY METAL'); row+=1
for k,v in collections.Counter(i['metal_detail'] for i in items).most_common(): kv(row,k,f'{v} items'); row+=1
row+=1
hdr(row,'BY CATEGORY'); row+=1
for k,v in collections.Counter(i['kind'] for i in items).most_common(): kv(row,k,f'{v} items'); row+=1

wb.save('Lab_Diamond_Jewelry_ALL_PHOTOS.xlsx')
print("items:",len(items),"| photos placed:",placed,"| missing:",missing,"| photo cols:",MAXP)
print("size: %.1f MB"%(os.path.getsize('Lab_Diamond_Jewelry_ALL_PHOTOS.xlsx')/1e6))
